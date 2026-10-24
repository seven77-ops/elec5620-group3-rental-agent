"""Build the offline snapshot from verified public research workbooks.

Run with the bundled Python/openpyxl extraction runtime, not the web server.
"""
import argparse
from collections import defaultdict
import hashlib
import json
from pathlib import Path
from statistics import quantiles
import openpyxl


def build(directory):
    manifest = json.loads((directory / 'sources.json').read_text())
    sources = [s for s in manifest['sources'] if s['id'] in {'nsw_bonds_2026_09', 'nsw_quarter_2026_06'}]
    for source in sources:
        assert hashlib.sha256((directory / source['local_file']).read_bytes()).hexdigest() == source['sha256']
    monthly_source, quarterly_source = sources
    wb = openpyxl.load_workbook(directory / monthly_source['local_file'], read_only=True, data_only=True)
    groups = defaultdict(list)
    for day, postcode, kind, beds, rent in wb.worksheets[0].iter_rows(min_row=4, values_only=True):
        assert (day.year, day.month) == (2026, 9)
        if postcode in (2008, 2037, 2042, 2050) and kind == 'F' and str(beds) in ('1', '2'):
            try:
                value = float(rent)
            except (TypeError, ValueError):
                continue
            if 0 < value < float('inf'):
                groups[(str(postcode), int(beds))].append(value)
    wb.close()
    cohorts = []
    for (postcode, beds), values in sorted(groups.items()):
        assert len(values) > 10
        q25, median, q75 = quantiles(values, n=4, method='inclusive')
        cohorts.append(dict(postcode=postcode, bedrooms=beds, period='2026-09', sample_size=len(values),
                            q25=q25, median=median, q75=q75, official_statistic=False,
                            source_id=monthly_source['id']))
    wb = openpyxl.load_workbook(directory / quarterly_source['local_file'], read_only=True, data_only=True)
    for index, row in enumerate(wb['Postcode'].iter_rows(min_row=10, values_only=True), 10):
        if row[0] in (2008, 2037, 2042, 2050) and row[1] == 'Flat/Unit' and row[2] in ('1 Bedroom', '2 Bedrooms'):
            cohorts.append(dict(postcode=str(row[0]), bedrooms=int(row[2][0]), period='2026-Q2',
                                sample_size=int(row[6].replace(',', '')), q25=float(row[3]),
                                median=float(row[4]), q75=float(row[5]), official_statistic=True,
                                source_id=quarterly_source['id'], source_range=f'Postcode!A{index}:L{index}'))
    wb.close()
    assert len(cohorts) == 16
    return {'snapshot_version': '1.0', 'compiled_on': '2026-10-21',
            'postcodes': ['2008', '2037', '2042', '2050'], 'bedrooms': [1, 2],
            'dwelling_type': 'flat_unit', 'rent_scope': 'whole_property', 'currency': 'AUD',
            'period_basis': 'Bond registration period; not current advertised availability',
            'quantile_method': 'Monthly: inclusive linear interpolation, not the official Rent Check calculation',
            'sources': [{k: s[k] for k in ('id', 'title', 'url', 'download_url', 'sha256')} for s in sources],
            'cohorts': cohorts}


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--research-dir', type=Path, required=True)
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    result = build(args.research_dir)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(result, ensure_ascii=False, indent=2, allow_nan=False) + '\n')
    print(f'Built {len(result["cohorts"])} cohorts with verified source hashes')
