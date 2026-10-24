from pathlib import Path
import json,re
from docx import Document
from docx.shared import Inches,Pt,RGBColor,Cm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'docs/user-scenarios/2026-10-22'
results=json.loads((OUT/'simulation-results.json').read_text())
doc=Document();sec=doc.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=Cm(1.8);sec.bottom_margin=Cm(1.8);sec.left_margin=Cm(1.9);sec.right_margin=Cm(1.9)
sec.footer_distance=Cm(.8)
for name,size,bold in [('Normal',10.5,False),('Title',22,True),('Subtitle',11,False),('Heading 1',16,True),('Heading 2',12,True),('Heading 3',10.5,True)]:
 st=doc.styles[name];st.font.name='Arial Unicode MS';st.font.size=Pt(size);st.font.bold=bold;st.font.italic=False;st.font.color.rgb=RGBColor(0,0,0)
 st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Arial Unicode MS')
 st.paragraph_format.space_after=Pt(5 if name=='Normal' else 8)
 st.paragraph_format.space_before=Pt(0 if name in ('Normal','Title','Subtitle') else 10)
 st.paragraph_format.line_spacing=Pt(15 if name=='Normal' else size*1.3)
 snap=OxmlElement('w:snapToGrid');snap.set(qn('w:val'),'0');st.element.get_or_add_pPr().append(snap)
 if name.startswith('Heading'):st.paragraph_format.keep_with_next=True
normal=doc.styles['Normal'];normal.paragraph_format.widow_control=True
# Stable, readable page numbering for the long test record.
f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.CENTER
run=f.add_run('第 ');run.font.size=Pt(9)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
f.add_run(' 页').font.size=Pt(9)
doc.core_properties.title='租房助手模拟用户场景与问答记录'
doc.core_properties.author='租房助手项目组'
doc.core_properties.subject='六类虚构用户的真实模型问答与保存结果'

md=[]
def p(text='',boldprefix=None,style=None):
 para=doc.add_paragraph(style=style)
 if boldprefix and text.startswith(boldprefix):para.add_run(boldprefix).bold=True;para.add_run(text[len(boldprefix):])
 else:para.add_run(text)
 if boldprefix=='用户：':para.paragraph_format.keep_with_next=True
 md.append(text+'\n')
 return para

def heading(text,level=1):doc.add_heading(text,level);md.append('#'*(level+1)+' '+text+'\n')

def bullet(text):
 para=p('• '+text);para.paragraph_format.left_indent=Cm(.25);para.paragraph_format.first_line_indent=Cm(-.25)
 return para

def table(headers,rows,widths):
 t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 for i,(h,w) in enumerate(zip(headers,widths)):t.columns[i].width=Cm(w);t.rows[0].cells[i].text=h
 for row in rows:
  cells=t.add_row().cells
  for i,value in enumerate(row):cells[i].text=str(value)
 for j,row in enumerate(t.rows):
  trpr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trpr.append(cant)
  if j==0:
   repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
  for i,cell in enumerate(row.cells):
   cell.width=Cm(widths[i]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   tcpr=cell._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
   for side in ('top','left','bottom','right'):
    el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
   tcpr.append(borders)
   margin=OxmlElement('w:tcMar')
   for side,value in [('top',95),('bottom',95),('left',100),('right',100)]:
    el=OxmlElement('w:'+side);el.set(qn('w:w'),str(value));el.set(qn('w:type'),'dxa');margin.append(el)
   tcpr.append(margin)
   if j==0:
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E8EFF8');tcpr.append(sh)
   for para in cell.paragraphs:
    para.paragraph_format.space_after=Pt(0);para.paragraph_format.space_before=Pt(0);para.paragraph_format.line_spacing=Pt(14)
    if i==0:para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for run in para.runs:run.font.size=Pt(9.5);run.bold=(j==0)
 md.append('| '+' | '.join(headers)+' |\n| '+' | '.join('---' for _ in headers)+' |\n')
 for row in rows:md.append('| '+' | '.join(str(x).replace('\n',' ') for x in row)+' |\n')
 doc.add_paragraph().paragraph_format.space_after=Pt(0)
 return t

verdict={
 'U01':('需优化','用户已说卧卫车位不限，AI仍重复追问，需求未完成确认。'),
 'U02':('部分通过','无结果、扩大区域和价格比较有效；单独提高预算只产生预览。'),
 'U03':('部分通过','草稿保存与资料接入可用；签约状态追问被误识别为候选跟进。'),
 'U04':('通过','供应商确认、待确认草稿改写、同一报修跟进和否定语句均符合预期。'),
 'U05':('部分通过','真实公司目录可查；价格和押金保证追问出现无关回答。'),
 'U06':('通过','指代不明时澄清，明确问题后只修改门锁记录，厨房问题保留。')}
summary_notes={
'U01':[
 '第1轮正确追问500澳元的周期与计价对象，没有直接把它作为整套周租。',
 '第2轮用户明确说卧室、浴室和车位均不限，AI仍在第2、3轮重复询问这三项；最终没有确认偏好或生成候选。',
 '第4轮没有声称广告已核实包账单，但回答混入噪音、隔音和通勤等无关内容，针对性不足。'],
'U02':[
 '第1轮Rhodes、2卧2卫、最高800澳元筛出0条，给出最低广告周租890及扩大区域的备选方向；没有编造匹配房源。',
 '用户确认第1轮条件后保存偏好v1。第2轮900澳元预览找到1条，但没有提供新的确认动作，已确认偏好当时仍为800。',
 '第3轮Parramatta条件预览找到3条，用户确认后保存偏好v2。第4轮比较650与700澳元广告，差额50；没有将其当作个人总费用。',
 '第5轮承认缺少已核实噪音及通勤数据，没有给出虚构排名。'],
'U03':[
 '前置候选为39/68 MacArthur Street, Parramatta。候选只是广告快照，尚未验证在租或拟签条款。',
 '合同助手的表单确认保存了v1未签署草稿；PDF接口成功返回且内容为PDF。拟定800澳元周租和3200押金是本场景虚构输入，不是广告提取或真实双方约定。',
 '第3轮询问草稿是否等于签约，AI却返回询问中介清单，没有直接解释当前状态。数据层仍未创建已入住或退租押金案件。',
 '模拟用户另行确认来源属于实际租约后，带入地址、代理姓名、邮箱及租约起止日期；供水电网选择仍为空，没有补造供应商。'],
'U04':[
 '用户确认后保存电力AGL、网络TPG；第2轮能读取同一份供应商记录。',
 '第3轮暂不确认报修草稿，第4轮礼貌改写后再确认，最终只创建1条报修。',
 '中介回复作为用户原话记录；“还没修好，不要标记解决”后状态仍为已联系。',
 '明确说已修好并确认后，原报修变为已解决；记录保留3条后续更新，邮件始终标记未发送。'],
'U05':[
 '确认后保存Rhodes、2卧2卫、有地毯、有电梯、2026-12-15预计搬出、目的地Burwood。',
 '第2轮返回3家搬运公司，第3轮返回3家清洁公司，卡片带目录中的官方网站和可用联系方式。没有创建预约或生成虚构报价。',
 '推荐文字仍使用通用询价清单，没有充分复用已保存的卧卫、地毯、日期和目的地。',
 '第4轮问哪家最便宜及能否保证押金，AI却谈噪音、隔音和通勤，属于答非所问。'],
'U06':[
 '前置保存厨房漏水和门锁损坏两条报修，均未解决，且没有对话聚焦记录。',
 '第1轮“那个问题”没有触发更新，AI要求说明具体报修。',
 '第2轮同句描述门锁已修好、厨房未修好，确认后只将门锁置为已解决。',
 '第3轮再次明确门锁后，仍只更新同一记录。最终厨房保持草稿待联系，门锁已解决，两条记录未合并。']}

p('租房助手模拟用户场景与问答记录',style='Title')
p('六类用户从找房到租后服务的实际调用记录',style='Subtitle')
p('记录日期 2026年10月22日')
p('本次用6个虚构用户完成29轮实际问答，覆盖需求澄清、房源比较、合同草稿、入住资料、报修、搬运与清洁。连续报修和维修指代区分表现正常；需求“不限”的识别、预算修改后的确认入口、合同状态追问及服务价格追问仍需改进。')
p('文档中的“用户”和“AI”逐轮保留本次实际文本；“用户操作”记录另行确认及保存结果。“建议回答”是后续优化示例，不代表系统本次已经这样回答。')
heading('模拟用户一览')
rows=[]
for r in results:
 a=r['persona'];rows.append([a['id'],a['name']+'\n'+a['role'],a['requirements'],verdict[a['id']][0]])
table(['编号','用户','主要要求','观察结果'],rows,[1.2,3.2,10.1,2.7])
heading('执行方式与结果口径')
p('使用项目当前后端及已配置的qwen-plus模型，按用户顺序发送消息。每轮调用单独同意外部AI；用户点击确认等操作通过同一后端接口复现。本轮是实际模型和接口场景模拟，不等同逐页浏览器操作测试。')
p('29轮请求均成功返回，29次问答前后检查均未发现租后资料被自动写入；另记录14项用户操作，包含前置记录和独立确认。接口返回成功不等于问答质量通过，逐场景评价以下列实际表现为准。')
p('所有用户和合同双方身份均为虚构，联系邮箱使用example.com。房源来自项目现有广告快照，价格、在租状态和计价对象仍待核实。测试在临时案例运行，不改原案例，不发送邮件或提交外部预约。')

for result in results:
 persona=result['persona'];pid=persona['id']
 doc.add_page_break();md.append('\n---\n')
 heading(pid+' '+persona['name']+' '+persona['role'])
 p('使用阶段：'+persona['stage'],boldprefix='使用阶段：')
 p('用户要求：'+persona['requirements'],boldprefix='用户要求：')
 p('本场景关注：'+'；'.join(persona['expectations'])+'。',boldprefix='本场景关注：')
 heading('实际问答与用户操作',2)
 for e in result['entries']:
  if e['type']=='dialogue':
   heading('第'+str(e['round'])+'轮',3)
   p('用户：'+e['user'],boldprefix='用户：')
   p('AI：'+e.get('answer',str(e.get('error','调用失败'))),boldprefix='AI：')
   if e.get('listings'):p('房源卡片：本轮展示'+str(len(e['listings']))+'条广告。')
   if e.get('providers') and persona['id']=='U05':
    p('公司卡片：'+'；'.join(x['name'] for x in e['providers'])+'。')
   acts=[x['label'] for x in e.get('actions',[]) if x['type'].startswith('apply_')]
   if acts:p('确认入口：'+'；'.join(acts)+'。仅对话没有执行保存。')
  elif e['type']=='user_action':
   label=e['label'];d=e['details']
   if label.startswith('前置状态'):
    detail=('候选地址 '+d['address']+'，在租状态尚未核实。') if d.get('address') else '厨房漏水与门锁损坏各1条，初始均待处理。'
   elif '偏好' in label:
    detail='确认后，需求版本为v'+str(e['state']['confirmation_versions'][-1])+'。'
   elif '合同助手' in label:
    detail='使用虚构双方资料，租客Alex Lin，代理Example Agent，邮箱agent@example.com；拟定每周800澳元、押金3200、租期2026-11-01至2027-11-01。保存草稿v1并取得PDF，未创建实际租约案件。'
   elif '接入住处' in label:
    detail='选择带入地址、代理姓名、邮箱和租约起止日期；租后资料保存为v'+str(e['state']['living_version'])+'。这一模拟确认不表示真实签约。'
   else:
    summary=d.get('summary',[])
    # Summarize confirmation effects without repeating the complete proposal text.
    v=e['state'];detail='确认保存后，租后资料为v'+str(v['living_version'])+'。'
    if v['repairs']:
     rr=v['repairs'][-1];status={'prepared':'草稿待联系','contacted':'已联系','resolved':'已解决'}[rr['status']]
     detail+='现有报修'+str(len(v['repairs']))+'条，本次对应记录为'+status+'，邮件未发送。'
    elif v['profile']['utilities'].get('electricity'):detail+='已保存AGL电力和TPG网络。'
    elif pid=='U05':detail+='已保存搬家日期、区域、卧卫及地毯和电梯条件。'
   para=p('用户操作：'+label+'。'+detail,boldprefix='用户操作：');para.paragraph_format.space_before=Pt(3)
  else:p('操作记录：'+str(e.get('error','异常')))
 heading('本场景结论',2)
 p('观察结果：'+verdict[pid][0]+'。'+verdict[pid][1],boldprefix='观察结果：')
 for note in summary_notes[pid]:bullet(note)

# Actionable findings distinguish actual behavior from suggested behavior.
doc.add_page_break();md.append('\n---\n')
heading('从模拟中发现的四项改进')
findings=[
 ('一 不限条件仍被追问','U01 第2至3轮',
  '用户已明确说卧室、浴室和车位数量都不限，但AI继续询问这些数量，导致本场景无法进入需求确认。',
  '把“不限”“无所谓”“不要求车位”等表达映射到对应条件，完整复述已收集信息，只追问真正缺少的项目。',
  '已记下：1人、Kensington、独立房间、最高每人每周500澳元，卧室、浴室和车位数量不限。广告计价对象和账单尚未核实，我可以先给你查看广告字段预览。'),
 ('二 预算预览与偏好保存衔接不足','U02 第2轮',
  '预算提高到900后查到1条广告，但本轮只有预览，没有新的确认入口，已确认需求仍暂时保留800。用户容易以为两者已经同步。',
  '明确区分当前预览与已保存需求；用户表达修改长期偏好时，提供“确认同步900并更新推荐”入口。',
  '这次预览已按最高每周900澳元查询，其他条件保持不变。你已保存的需求仍是800澳元；核对后可以确认同步为900并更新推荐。'),
 ('三 签约状态问题被误路由','U03 第3轮',
  '用户问“草稿生成是否等于签约和入住”，系统识别为候选跟进，返回了问中介的通用清单，没有直接回答当前状态。虽然没有错误创建租约，问答仍未解决用户疑惑。',
  '增加草稿、签约、实际入住的状态问答，将已保存草稿及实际租约状态用于路由，并直接解释下一步确认。',
  '当前保存的是未签署草稿，不能因此认定签约完成或已经入住。请先完成双方核对；确认实际租约后，再选择把住处、联系人和租期带入租后服务。'),
 ('四 清洁价格追问丢失服务语境','U05 第4轮',
  '前一轮在聊退租清洁，追问哪家最便宜、是否保证押金时，却返回缺少噪音、隔音、通勤数据的房源回答，属于明显答非所问。',
  '让澄清回复沿用当前服务主题。没有实时报价时解释这一限制，继续帮助整理报价项目，不套用房源通用提示。推荐时也应复用已保存的卧卫、地毯和日期。',
  '目前没有这几家公司的实时报价，我不能判断哪家肯定最便宜，也不能保证押金结果。你可以把收到的报价发来，我帮你比较地毯清洁、烤箱、窗户、额外收费和服务日期。')]
for title,source,observed,fix,suggestion in findings:
 heading(title,2);p('对应记录：'+source,boldprefix='对应记录：');p('实际表现：'+observed,boldprefix='实际表现：');p('优化方向：'+fix,boldprefix='优化方向：');p('建议回答：'+suggestion,boldprefix='建议回答：')
heading('下一轮可复用的验收问题')
table(['场景','验收重点'],[
 ['说都不限','不重复追问已明确的卧卫及车位要求，能够进入确认步骤。'],
 ['提高预算但其他不变','完整保留区域和卧卫，展示当前预览与已保存版本的区别。'],
 ['草稿是不是签约','直接回答当前草稿状态，不跳回询问中介清单，不自动认定入住。'],
 ['比较清洁公司报价','承接清洁语境，承认未知价格，提出具体可比较项目。'],
 ['还没修好不要解决','保持未解决，保存原话需要单独确认。'],
 ['那个问题好了','没有明确聚焦时询问哪条；明确问题后只修改对应记录。']], [4.0,13.2])
p('建议优先修正需求“不限”识别、合同状态问答和租后服务追问路由，再补预算修改的确认入口。随后用相同用户消息重跑，逐轮对照本记录，不以接口成功率代替问答质量。')
heading('复现材料',2)
p('同目录保留simulation-results.json及六个用户的逐条记录。可复现脚本为tools/simulate_users_20261022.py，需使用项目当前AI配置；每次运行使用临时案例，模型结果可能随版本或采样变化。')

# Remove theme title borders and font overrides from the default template.
for style in doc.styles:
 for border in list(style.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
 for fonts in style.element.iter(qn('w:rFonts')):
  for attr in ('asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme','csTheme'):
   fonts.attrib.pop(qn('w:'+attr),None)
  for attr in ('ascii','hAnsi','eastAsia','cs'):fonts.set(qn('w:'+attr),'Arial Unicode MS')
for para in doc.paragraphs:
 for border in list(para._p.iter(qn('w:pBdr'))):border.getparent().remove(border)
path=OUT/'租房助手模拟用户场景与问答记录.docx';doc.save(path)
(OUT/'租房助手模拟用户场景与问答记录.md').write_text('\n'.join(md))
print(path)
print('paragraphs',len(doc.paragraphs),'tables',len(doc.tables))
